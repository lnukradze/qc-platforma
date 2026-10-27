# -*- coding: utf-8 -*-
"""
T-015 — კონტროლი ველისთვის ანონიმურია; კოორდინატორი — ოპერატორების სტატისტიკის გარეშე.

ყველა სახელი გამოგონილია და ლათინურია, რომ ქართულ წარწერებს (როლები, ღილაკები) არ დაემთხვეს.
"""
import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import domain as D
from app import scope
from app.main import app
from app.models import EmailMessage, Presence, QcEvent, QcItem, User
from app.services import anon, notify
from app.workflow import apply
from conftest import ctx, make_item, make_user

ORIGIN = {"origin": "http://testserver"}
FIELD = [("sv", D.SURVEYOR), ("mgr", D.FIELD_MANAGER)]


def login(c, user, role):
    r = c.post("/login/demo", data={"user_id": str(user.id), "role": role}, headers=ORIGIN,
               follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers.get("location")


def client(user, role):
    c = TestClient(app)
    login(c, user, role)
    return c


@pytest.fixture(autouse=True)
def no_smtp(monkeypatch):
    """რეალური წერილი არ იგზავნება: SMTP გამორთულია, გაგზავნის ცდა ტესტს აგდებს."""
    monkeypatch.setattr(notify, "SMTP_HOST", "")

    def _boom(msg):
        raise AssertionError("SMTP-ზე გაგზავნა ტესტში არ უნდა მოხდეს")
    monkeypatch.setattr(notify, "_send", _boom)


@pytest.fixture()
def w(db):
    """გამოგონილი სამყარო: QC-ის ყველა როლი, ველის გუნდი, ხილული / გაუგზავნელი / სხვისი ოქმი."""
    p = {
        "op": make_user(db, "Tamar Qcop-Alpha", D.QC_OPERATOR),
        "op2": make_user(db, "Nino Qcop-Beta", D.QC_OPERATOR),
        "qcm": make_user(db, "Levan Qcm-Gamma", D.QC_MANAGER),
        "admin": make_user(db, "Irakli Admin-Delta", D.ADMIN),
        "mgr": make_user(db, "Giorgi Field-Primary", D.FIELD_MANAGER, email="mgr@example.com"),
        "mgr2": make_user(db, "Eka Field-Copy", D.FIELD_MANAGER, email="mgr2@example.com"),
        "sv": make_user(db, "Sopo Field-Surveyor", D.SURVEYOR),
        "sv2": make_user(db, "Dato Field-Second", D.SURVEYOR),
        "mgrx": make_user(db, "Keti Other-Manager", D.FIELD_MANAGER, email="mgrx@example.com"),
        "svx": make_user(db, "Beka Other-Surveyor", D.SURVEYOR),
        "coord": make_user(db, "Ana Coord-Eta", D.COORDINATOR),
        "multi": make_user(db, "Mari Multi-Theta", D.QC_OPERATOR, D.SURVEYOR),
    }
    db.commit()
    a = make_item(db, p, protocol="6200001", managers=("mgr", "mgr2"))
    u = make_item(db, p, protocol="6200002", cad="202.51.12.40.001.002")
    x = make_item(db, p, protocol="6200003", cad="202.51.12.40.001.003", managers=("mgrx",),
                  surveyors=("svx",))
    ev_send = apply(db, a, p["op"], D.QC_OPERATOR, "send",
                    {"version": a.version, "verdict": D.NEGATIVE, "qc_comment": "ფოტო აკლია"})
    apply(db, a, p["sv"], D.SURVEYOR, "field_respond", {"version": a.version, "response": D.R_DISAGREE, "comment": "ვერ ვეთანხმები"})
    apply(db, a, p["mgr"], D.FIELD_MANAGER, "manager_approve", {"version": a.version})
    ev_stand = apply(db, a, p["qcm"], D.QC_MANAGER, "qcm_stand", {"version": a.version, "comment": "ხარვეზი რჩება"})
    ev_admin = apply(db, a, p["admin"], D.ADMIN, "comment", {"once": "test-once-1", "comment": "გთხოვთ გაასწოროთ"})
    apply(db, a, p["op"], D.QC_OPERATOR, "request_verdict_change",
          {"version": a.version, "new_verdict": D.RECOMMENDATION, "comment": "რეკომენდაცია უფრო ზუსტია"})
    apply(db, x, p["op"], D.QC_OPERATOR, "send", {"version": x.version, "verdict": D.NEGATIVE, "qc_comment": "x"})
    u.operator_id = p["op2"].id
    now = dt.datetime.now()
    db.add_all([
        Presence(item_id=a.id, user_id=p["op2"].id, role=D.QC_OPERATOR, last_seen=now),
        Presence(item_id=a.id, user_id=p["qcm"].id, role=D.QC_MANAGER, last_seen=now),
        Presence(item_id=a.id, user_id=p["admin"].id, role=D.ADMIN,
                 last_seen=now - dt.timedelta(minutes=5)),                      # ვადაგასული
        Presence(item_id=a.id, user_id=p["mgr2"].id, role=D.FIELD_MANAGER, last_seen=now),
        Presence(item_id=x.id, user_id=p["op"].id, role=D.QC_OPERATOR, last_seen=now),  # უხილავი
        Presence(item_id=u.id, user_id=p["op2"].id, role=D.QC_OPERATOR, last_seen=now),  # გაუგზავნელი
    ])
    db.commit()
    qc = [p[k] for k in ("op", "op2", "qcm", "admin", "multi")]
    return {"p": p, "a": a, "u": u, "x": x, "qc": qc,
            "events": {"send": ev_send, "stand": ev_stand, "admin": ev_admin}}


def assert_no_qc_names(w, text, where=""):
    leaked = [u.full_name for u in w["qc"] if u.full_name in text]
    assert not leaked, (where, leaked)


# ── 1. ერთი წესი ─────────────────────────────────────────────────────────
@pytest.mark.parametrize("viewer", [D.SURVEYOR, D.FIELD_MANAGER])
@pytest.mark.parametrize("actor_role", [D.QC_OPERATOR, D.QC_MANAGER, D.ADMIN])
def test_actor_name_hides_qc_from_field(db, w, viewer, actor_role):
    who = {D.QC_OPERATOR: "op", D.QC_MANAGER: "qcm", D.ADMIN: "admin"}[actor_role]
    assert anon.actor_name(viewer, w["p"][who], actor_role) == D.QC_ANON_LABEL


@pytest.mark.parametrize("viewer", [D.QC_OPERATOR, D.QC_MANAGER, D.ADMIN, D.COORDINATOR])
def test_actor_name_qc_viewer_sees_names(db, w, viewer):
    assert anon.actor_name(viewer, w["p"]["op"], D.QC_OPERATOR) == "Tamar Qcop-Alpha"
    assert anon.actor_name(viewer, w["p"]["sv"], D.SURVEYOR) == "Sopo Field-Surveyor"


@pytest.mark.parametrize("viewer", [D.SURVEYOR, D.FIELD_MANAGER])
def test_actor_name_field_author_and_multi_role(db, w, viewer):
    p = w["p"]
    assert anon.actor_name(viewer, p["mgr"], D.FIELD_MANAGER) == "Giorgi Field-Primary"
    # მრავალროლიანი: კონკრეტული მოქმედების როლი წყვეტს
    assert anon.actor_name(viewer, p["multi"], D.SURVEYOR) == "Mari Multi-Theta"
    assert anon.actor_name(viewer, p["multi"], D.QC_OPERATOR) == D.QC_ANON_LABEL
    # როლი უცნობია (ძველი ჩანაწერი) — QC როლის მქონე სახელი არ ჩანს
    assert anon.actor_name(viewer, p["multi"], None) == D.QC_ANON_LABEL
    assert anon.actor_name(viewer, p["sv2"], None) == "Dato Field-Second"
    assert anon.actor_name(viewer, None, None) == ""
    assert anon.actor_name(viewer, None, D.QC_OPERATOR) == D.QC_ANON_LABEL
    assert p["multi"].full_name == "Mari Multi-Theta"             # ORM უცვლელია


def test_render_actor_name_follows_role_between_requests(db, w):
    """render()-ის callable — ამ მოთხოვნის როლით, სესიებს შორის არ გადადის."""
    p = w["p"]
    c = client(p["multi"], D.SURVEYOR)
    r = c.get("/stats")
    assert r.status_code == 200
    assert r.context["actor_name"](p["op"], D.QC_OPERATOR) == D.QC_ANON_LABEL
    other = client(p["qcm"], D.QC_MANAGER)
    r2 = other.get("/stats")
    assert r2.context["actor_name"](p["op"], D.QC_OPERATOR) == "Tamar Qcop-Alpha"
    login(c, p["multi"], D.QC_OPERATOR)                           # იგივე ადამიანი, სხვა როლი
    r3 = c.get("/stats")
    assert r3.context["actor_name"](p["op"], D.QC_OPERATOR) == "Tamar Qcop-Alpha"
    assert r.context["actor_name"](p["op"], D.QC_OPERATOR) == D.QC_ANON_LABEL   # ძველი არ შეიცვალა


# ── 2–3. API და presence — ველის ორივე როლი ──────────────────────────────
@pytest.mark.parametrize("who,role", FIELD)
def test_api_items_field_no_operator(db, w, who, role):
    c = client(w["p"][who], role)
    r = c.get("/api/items")
    assert r.status_code == 200
    rows = r.json()["items"]
    assert [i["id"] for i in rows] == [w["a"].id]                 # გაუგზავნელი და სხვისი — არა
    assert all("operator" not in i for i in rows)
    assert_no_qc_names(w, r.text, "/api/items")
    r = c.get(f"/api/items/{w['a'].id}")
    assert r.status_code == 200
    d = r.json()
    assert "operator" not in d and d["history"]
    assert_no_qc_names(w, r.text, "/api/items/{id}")
    for hidden in (w["u"], w["x"]):
        assert c.get(f"/api/items/{hidden.id}").status_code == 404


@pytest.mark.parametrize("who,role", FIELD)
def test_presence_field_anonymous(db, w, who, role):
    c = client(w["p"][who], role)
    a, u, x = w["a"].id, w["u"].id, w["x"].id
    r = c.get(f"/api/presence?ids={a},{u},{x}")
    assert r.status_code == 200
    d = r.json()
    assert set(d) == {str(a)}                                     # უხილავი ოქმი — არა
    people = d[str(a)]
    qc = [o for o in people if o["name"] == D.QC_ANON_LABEL]
    assert len(qc) == 2                                           # ორი აქტიური QC; ვადაგასული — არა
    assert all("role" not in o for o in qc)
    assert {"name": "Eka Field-Copy", "role": D.ROLE_SHORT[D.FIELD_MANAGER]} in people
    assert_no_qc_names(w, r.text, "GET /api/presence")
    r = c.post(f"/api/presence/{a}", headers=ORIGIN)
    assert r.status_code == 200
    qc = [o for o in r.json()["others"] if o["name"] == D.QC_ANON_LABEL]
    assert len(qc) == 2 and all("role" not in o for o in qc)
    assert_no_qc_names(w, r.text, "POST /api/presence")
    assert c.post(f"/api/presence/{x}", headers=ORIGIN).status_code == 404
    assert c.post(f"/api/presence/{a}/leave", headers=ORIGIN).json() == {"ok": True}


def test_presence_qc_viewer_sees_names(db, w):
    c = client(w["p"]["qcm"], D.QC_MANAGER)
    people = c.get(f"/api/presence?ids={w['a'].id}").json()[str(w["a"].id)]
    names = {x["name"]: x.get("role") for x in people}
    assert names["Nino Qcop-Beta"] == D.ROLE_SHORT[D.QC_OPERATOR]
    assert "Eka Field-Copy" in names and "Irakli Admin-Delta" not in names   # ვადაგასული


# ── 4. მეილი ─────────────────────────────────────────────────────────────
@pytest.mark.parametrize("key,sender,kind", [("send", "op", "sent"), ("stand", "qcm", "qcm_stand"),
                                             ("admin", "admin", "comment")])
def test_email_sender_anonymous_for_field(db, w, key, sender, kind):
    ev = db.get(QcEvent, w["events"][key].id)
    item = db.get(QcItem, w["a"].id)
    msg = notify.queue(db, item, ev, w["p"][sender], kind, base_url="http://x/")
    db.commit()
    assert msg.status == "preview"
    assert msg.to_user_id == w["p"]["mgr"].id and msg.cc == ["mgr2@example.com"]   # ორივე ველის
    for body in (msg.text, msg.html):
        assert D.QC_ANON_LABEL in body
        assert_no_qc_names(w, body, kind)
    assert "Giorgi Field-Primary" in msg.text and "Giorgi Field-Primary" in msg.html
    assert msg.created_by_id == w["p"][sender].id                  # აუდიტი — რეალური ავტორი


def test_email_via_http_action(db, w):
    c = client(w["p"]["qcm"], D.QC_MANAGER)
    r = c.post(f"/items/{w['a'].id}/action", headers=ORIGIN, follow_redirects=False,
               data={"action": "comment", "comment": "შეამოწმეთ ხელახლა", "notify": "on",
                     "once": "test-once-http"})
    assert r.status_code == 303 and "err=" not in r.headers["location"]
    msg = db.query(EmailMessage).order_by(EmailMessage.id.desc()).first()
    assert msg is not None and msg.created_by_id == w["p"]["qcm"].id
    for body in (msg.text, msg.html):
        assert D.QC_ANON_LABEL in body
        assert_no_qc_names(w, body, "http")


def test_email_sender_named_for_qc_only_recipients(db, w):
    """ადრესატი QC-ის მხარეს არის (ველის როლი არ აქვს) — სახელი რჩება."""
    ev = db.get(QcEvent, w["events"]["send"].id)
    assert notify.sender_name(ev, w["p"]["op"], [w["p"]["qcm"]]) == "Tamar Qcop-Alpha"
    assert notify.sender_name(ev, w["p"]["op"], [w["p"]["qcm"], w["p"]["mgr2"]]) == D.QC_ANON_LABEL


# ── 5. ბარათის კონტექსტი (backend) ───────────────────────────────────────
@pytest.mark.parametrize("who,role", FIELD)
def test_detail_context_field(db, w, who, role):
    c = client(w["p"][who], role)
    r = c.get(f"/items/{w['a'].id}")
    assert r.status_code == 200
    assert r.context["show_operator"] is False and r.context["operators"] == []
    fn = r.context["actor_name"]
    req = r.context["pending_req"]
    if req is not None:                    # მოთხოვნას როლი არ აქვს — ავტორის როლებით (სესია დახურულია)
        assert fn(db.get(User, req.requested_by_id), None) == D.QC_ANON_LABEL
    for e in r.context["events"]:
        if e.actor_role in D.QC_ROLES:
            assert fn(e.actor, e.actor_role) == D.QC_ANON_LABEL
    field_ev = [e for e in r.context["events"] if e.actor_role == D.SURVEYOR]
    assert field_ev and fn(field_ev[0].actor, D.SURVEYOR) == "Sopo Field-Surveyor"


def test_detail_context_qc_manager(db, w):
    c = client(w["p"]["qcm"], D.QC_MANAGER)
    r = c.get(f"/items/{w['a'].id}")
    assert r.status_code == 200 and r.context["show_operator"] is True
    assert {n for _i, n in r.context["operators"]} >= {"Tamar Qcop-Alpha", "Nino Qcop-Beta"}
    assert "Sopo Field-Surveyor" in r.text and "Tamar Qcop-Alpha" in r.text


# ── 6. HTML ზედაპირები — შაბლონები T-016 / T-014-ში ────────────────────
@pytest.mark.xfail(strict=True, reason="T-016: items/detail.html ჯერ item.operator.full_name / "
                   "e.actor.full_name / pending_req.requested_by-ს ბეჭდავს — show_operator და actor_name გამოიყენოს")
@pytest.mark.parametrize("who,role", FIELD)
def test_detail_html_field_no_qc_names(db, w, who, role):
    r = client(w["p"][who], role).get(f"/items/{w['a'].id}")
    assert r.status_code == 200
    assert_no_qc_names(w, r.text, "/items/{id}")


@pytest.mark.xfail(strict=True, reason="T-014: home.html „ბოლო მოქმედებები“ ჯერ e.actor.full_name-ს ბეჭდავს "
                   "— actor_name გამოიყენოს")
@pytest.mark.parametrize("who,role", FIELD)
def test_home_html_field_no_qc_names(db, w, who, role):
    r = client(w["p"][who], role).get("/")
    assert r.status_code == 200
    assert_no_qc_names(w, r.text, "/")


@pytest.mark.parametrize("who,role", FIELD)
def test_home_context_field(db, w, who, role):
    r = client(w["p"][who], role).get("/")
    assert r.status_code == 200
    assert "pending" not in r.context["data"]["kpi"]
    fn = r.context["actor_name"]
    qc_ev = [e for e in r.context["recent"] if e.actor_role in D.QC_ROLES]
    assert qc_ev and all(fn(e.actor, e.actor_role) == D.QC_ANON_LABEL for e in qc_ev)


# ── 7. სტატისტიკა ველისთვის ──────────────────────────────────────────────
@pytest.mark.parametrize("who,role", FIELD)
def test_stats_field_no_operators(db, w, who, role):
    c = client(w["p"][who], role)
    r = c.get("/stats")
    assert r.status_code == 200 and r.context["opts"]["operators"] == []
    assert_no_qc_names(w, r.text, "/stats")
    d = c.get("/api/stats").json()
    assert "operators" not in d and "operators" not in d["sections"]
    assert "pending" not in d["kpi"] and d["kpi"]["checked"] == 1
    assert_no_qc_names(w, c.get("/api/stats").text, "/api/stats")
    # ხელით მიწოდებული ოპერატორის ფილტრი ჭრილს არ ქმნის
    for op in ("op", "op2"):
        f = c.get(f"/api/stats?operator={w['p'][op].id}").json()
        assert f["kpi"]["checked"] == d["kpi"]["checked"] and "operators" not in f
    wb = load_workbook(io.BytesIO(c.get("/stats/export.xlsx").content))
    assert "ოპერატორები" not in wb.sheetnames
    labels = [row[0].value for row in wb["ჯამური"].iter_rows(min_row=2)]
    assert "გასავლელი" not in labels and "დადებითი" in labels
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            assert_no_qc_names(w, " ".join(str(v) for v in row if v is not None), ws.title)


def test_stats_qc_manager_operator_filter_works(db, w):
    c = client(w["p"]["qcm"], D.QC_MANAGER)
    r = c.get("/stats")
    assert {o["name"] for o in r.context["opts"]["operators"]} == {"Tamar Qcop-Alpha", "Nino Qcop-Beta"}
    d = c.get("/api/stats").json()
    assert "pending" in d["kpi"] and {o["name"] for o in d["operators"]} >= {"Tamar Qcop-Alpha"}
    assert d["kpi"]["checked"] == 2
    f = c.get(f"/api/stats?operator={w['p']['op2'].id}").json()
    assert f["kpi"]["checked"] == 0


# „გასავლელი" — ერთი წესი (scope.can_see_pending) KPI-ში და Excel-ში, ყველა როლზე.
# w-ში გაუგზავნელია მხოლოდ u (ოპერატორი op2): op-ს თავისი 0 აქვს, op2-ს — 1, ყველას ნახვით — 1.
PENDING = [("sv", D.SURVEYOR, None), ("mgr", D.FIELD_MANAGER, None), ("head", D.HEAD_OF_SERVICE, None),
           ("op", D.QC_OPERATOR, 0), ("op2", D.QC_OPERATOR, 1), ("qcm", D.QC_MANAGER, 1),
           ("admin", D.ADMIN, 1), ("coord", D.COORDINATOR, 1)]


@pytest.mark.parametrize("who,role,expected", PENDING)
def test_pending_kpi_by_role(db, w, who, role, expected):
    user = w["p"][who] if who != "head" else make_user(db, "Gela Head-Iota", D.HEAD_OF_SERVICE)
    db.commit()
    assert scope.can_see_pending(ctx(user, role)) is (expected is not None)
    c = client(user, role)
    k = c.get("/api/stats").json()["kpi"]
    if expected is None:
        assert "pending" not in k
    else:
        assert k["pending"] == expected
    wb = load_workbook(io.BytesIO(c.get("/stats/export.xlsx").content))
    rows = {r[0]: r[1] for r in wb["ჯამური"].iter_rows(min_row=2, values_only=True)}
    if expected is None:
        assert "გასავლელი" not in rows
    else:
        assert rows["გასავლელი"] == expected


# ── 8. კოორდინატორი ─────────────────────────────────────────────────────
def test_coordinator_items_and_stats(db, w):
    c = client(w["p"]["coord"], D.COORDINATOR)
    r = c.get("/api/items")
    assert r.status_code == 200
    assert {i["id"] for i in r.json()["items"]} == {w["a"].id, w["x"].id}     # გაუგზავნელი — არა
    assert c.get(f"/api/items/{w['u'].id}").status_code == 404
    d = c.get("/api/stats").json()
    assert "operators" not in d and "operators" not in d["sections"] and "managers" in d
    assert c.get("/stats").context["opts"]["operators"] == []
    f = c.get(f"/api/stats?operator={w['p']['op2'].id}").json()
    assert f["kpi"]["checked"] == d["kpi"]["checked"]
    wb = load_workbook(io.BytesIO(c.get("/stats/export.xlsx").content))
    assert "ოპერატორები" not in wb.sheetnames and "მენეჯერები" in wb.sheetnames
    labels = [row[0].value for row in wb["ჯამური"].iter_rows(min_row=2)]
    assert "დადებითი" in labels and "დადებითი (+დაზუსტება)" not in labels


def test_coordinator_view_only(db, w):
    c = client(w["p"]["coord"], D.COORDINATOR)
    before = db.query(QcEvent).filter(QcEvent.item_id == w["a"].id).count()
    r = c.post(f"/items/{w['a'].id}/action", headers=ORIGIN, follow_redirects=False,
               data={"action": "comment", "comment": "x"})
    assert r.status_code in (303, 403) and (r.status_code == 403 or "err=" in r.headers["location"])
    db.expire_all()
    assert db.query(QcEvent).filter(QcEvent.item_id == w["a"].id).count() == before


def test_qc_manager_export_keeps_operators(db, w):
    c = client(w["p"]["qcm"], D.QC_MANAGER)
    wb = load_workbook(io.BytesIO(c.get("/stats/export.xlsx").content))
    assert "ოპერატორები" in wb.sheetnames
    assert "დადებითი" in [row[0].value for row in wb["ჯამური"].iter_rows(min_row=2)]


# ── 9. უარყოფითი კონტროლი — ანონიმურობა ერთმხრივია ──────────────────────
def test_qc_manager_sees_field_names(db, w):
    c = client(w["p"]["qcm"], D.QC_MANAGER)
    d = c.get(f"/api/items/{w['a'].id}").json()
    assert d["operator"] == "Tamar Qcop-Alpha" and d["field_manager"] == "Giorgi Field-Primary"
    assert "Sopo Field-Surveyor" in d["surveyors"]
    assert "Giorgi Field-Primary" in c.get("/api/stats").text
    assert "Sopo Field-Surveyor" in c.get(f"/items/{w['a'].id}").text
