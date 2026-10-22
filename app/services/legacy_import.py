# -*- coding: utf-8 -*-
"""
ძველი კონტროლის ცხრილის გადმოტანა → `qc_items`.

წყარო: `data/raw/control/*.json` (fetch_sheets.py, მხოლოდ კითხვით წამოღებული).

წესები (მომხმარებლის პასუხები, 2026-10-19 / 2026-10-20):
  • მთავარი წყარო — კონტროლის ცხრილი; დუბლიკატი კოდები არ უნდა გადმოვიდეს.
    თუ კოდი სპორადულშიცაა — რჩება მხოლოდ სპორადულში;
  • ვისთანაა „ბურთი" — Q (კონტროლის სინქრონი) და S (ველის სინქრონი) დროებით:
    Q > S და O-ში შენიშვნაა → კონტროლმა უკვე უპასუხა → საველე მენეჯერთან;
  • D — ოპერატორის მიერ სტატუსის მინიჭების დრო (რეპორტები მისით ითვლება);
  • M — საბოლოო დადასტურება (მხოლოდ უარყოფით/რეკომენდაციაზე);
  • N — მოძველებული, არ გამოიყენება; L „წაიშალა" — მენეჯერმა კომენტარი წაშალა;
  • K-ში ზოგჯერ ახალი ოქმის ნომერია; V/W (TYPE/anot) — ნედლად ვინახავთ.

⚠️ ძველი ცხრილი სტატუსებს ზედ წერდა, ამიტომ **ისტორია აღდგენადი არ არის** —
   ვინახავთ ბოლო მდგომარეობას + ნედლ ანაბეჭდს, დროები მიახლოებითია (approx).
"""
import datetime as dt
import json
import re
import urllib.parse
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from sqlalchemy import delete, exists, select

from .. import domain as D
from ..config import QC_WORKBOOK, RAW_DIR
from ..models import (ImportRun, ItemPerson, PersonAlias, QcEvent, QcItem, QcItemCode, QcReview,
                      User)
from .codes import parse_codes
from .distribution.algo import block_key, sector_of, zone_of
from .people import PID_RE, clean_name, name_key, split_names
from .seed import People

SPORADIC_TAB = "სპორადული"

HEADER_MAP = {
    "ოქმის N": "proto", "CADCODE": "cad", "QC სპეციალისტი": "qc",
    "დამატების თარიღი": "added", "შემოწმების სტატუსი": "status",
    "ამგეგმავი": "p1", "ამგეგმავი (2)": "p2",
    "ხარვეზის კოდი": "codes", "ხარვეზის ტიპი": "codes",
    "QC კომენტარი": "comment",
    "საველე მენეჯერის სტატუსი": "jstatus",
    "საველე მენეჯერის კომენტარი": "kcomment", "საველე ჯგუფის კომენტარი": "kcomment",
    "უკუკავშირის თარიღი": "ldate", "გასწორდა": "fixed", "ხარვეზია": "isdefect",
    "უკან დაბრუნებული ხარვეზის კომენტარი": "returned",
    "კონტროლი sync": "psync", "საველე sync": "rsync",
    "მენეჯერი": "manager", "ნახაზის შეტანის თარიღი": "drawing",
    "TYPE": "type", "anot": "anot",
}
SYNC_TIME = "სინქრონის დრო"          # ორჯერ გვხვდება: პირველი — Q, მეორე — S

_DT_FORMATS = ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
               "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%d.%m.%y")


def parse_dt(v):
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in _DT_FORMATS:
        try:
            return dt.datetime.strptime(s[:19] if "T" in s else s, fmt)
        except ValueError:
            continue
    return None


def parse_bool(v):
    return str(v or "").strip().lower() in ("true", "1", "კი", "დიახ")


def row_to_rec(header, row):
    rec, extra, sync_seen = {}, {}, 0
    for i, h in enumerate(header):
        v = row[i] if i < len(row) else ""
        v = "" if v is None else str(v)
        h = str(h or "").strip()
        if h == SYNC_TIME:
            rec["qtime" if sync_seen == 0 else "stime"] = v
            sync_seen += 1
        elif h in HEADER_MAP:
            rec.setdefault(HEADER_MAP[h], v)
        elif v.strip():
            extra[h or f"col{i + 1}"] = v
    rec["extra"] = extra
    return rec


def _score(r):
    s = 0.0
    if r.get("status", "").strip():
        s += 4
    if r.get("jstatus", "").strip():
        s += 2
    if parse_bool(r.get("fixed")):
        s += 1
    s += 0.01 * sum(1 for k, v in r.items() if k != "extra" and str(v).strip())
    d = parse_dt(r.get("added"))
    return (s, d or dt.datetime.min)


def qc_answered_last(r) -> bool:
    """კონტროლმა ველის ბოლო პასუხს უკვე უპასუხა: Q > S და O-ში შენიშვნაა (პასუხი №15)."""
    q, s_ = parse_dt(r.get("qtime")), parse_dt(r.get("stime"))
    return bool(q and s_ and q > s_ and r.get("returned", "").strip())


def derive_state(r):
    """ძველი მწკრივი → (state, close_reason, verdict). წესები ARCHITECTURE.md-შია აღწერილი."""
    verdict = D.VERDICT_FROM_SHEET.get(r.get("status", "").strip())
    resp = D.RESPONSE_FROM_SHEET.get(r.get("jstatus", "").strip())
    fixed = parse_bool(r.get("fixed"))
    if verdict is None:                                   # ცარიელი / „პროცესშია"
        return D.PENDING, None, None
    if verdict == D.POSITIVE:
        return D.CLOSED, D.CR_POSITIVE, verdict
    if fixed:
        return D.CLOSED, (D.CR_CORRECTED if verdict == D.NEGATIVE else D.CR_ACKNOWLEDGED), verdict
    if resp == D.R_DISAGREE:
        # უკვე მოსული არ დათანხმებული → QC მენეჯერთან; თუ მას უკვე უპასუხია → მენეჯერთან
        return (D.RETURNED_TO_MANAGER if qc_answered_last(r) else D.AWAITING_QCM), None, verdict
    if resp in (D.R_CORRECTED, D.R_ACKNOWLEDGE):
        # QC-მ უკვე დააბრუნა („არ გასწორდა") → ველზე; თუ არა — QC-ის გადასამოწმებელია
        return (D.RETURNED_TO_MANAGER if qc_answered_last(r) else D.AWAITING_QC), None, verdict
    if r.get("returned", "").strip():
        return D.RETURNED_TO_MANAGER, None, verdict
    return D.AWAITING_FIELD, None, verdict


def load_planner_aliases(db, ppl):
    """QC სამუშაო.xlsx → Reference: ნედლი სახელი → სტანდარტული (+ პირადი ნომერი)."""
    if not Path(QC_WORKBOOK).exists():
        return 0
    wb = openpyxl.load_workbook(QC_WORKBOOK, read_only=True, data_only=True)
    if "Reference" not in wb.sheetnames:
        return 0
    n = 0
    existing = {a.alias_key for a in db.scalars(select(PersonAlias))}
    for row in wb["Reference"].iter_rows(min_row=6, values_only=True):
        raw, std, pid = (row[0], row[1], row[2]) if len(row) > 2 else (None, None, None)
        if not raw or not std:
            continue
        u = ppl.get(std)
        if u is None:
            continue
        pid = str(pid or "").strip()
        if pid and PID_RE.fullmatch(pid) and not u.personal_id:
            u.personal_id = pid                           # 🔐 ინახება, ინტერფეისში არ ჩანს
        k = name_key(raw)
        if k and k != u.name_key and k not in existing:
            db.add(PersonAlias(alias_key=k, user_id=u.id, source="reference"))
            existing.add(k)
            ppl.by_key.setdefault(k, u)
            n += 1
    wb.close()
    db.flush()
    return n


def fill_people_from_source(db) -> int:
    """
    ცხრილში ამზომველი არ ეწერა (უმეტესად გაუვლელი ოქმები) → REGFIELDALL-იდან, იმავე წესით,
    რაც ახალ განაწილებაში: OQMI/MOXAZ, FIELD_GROUP_ALL-თან შედარებით (მხოლოდ კითხვა).
    """
    from .directory import FieldDirectory
    from .distribution import source
    missing = db.execute(select(QcItem.id, QcItem.protocol_number, QcItem.cadcode).where(
        QcItem.is_legacy.is_(True),
        ~exists().where(ItemPerson.item_id == QcItem.id, ItemPerson.kind == D.P_SURVEYOR))).all()
    if not missing:
        return 0
    try:
        found = source.lookup_protocols([p for _, p, _ in missing])
    except FileNotFoundError:
        return 0
    directory = FieldDirectory(db)
    n = 0
    for iid, proto, cad in missing:
        hits = found.get(proto) or []
        hit = next((h for h in hits if h["cadcode"] == cad), hits[0] if hits else None)
        if hit is None:
            continue
        names = [hit["moxaz"]] if " ".join(hit["oqmi"].split()) == "სისტემური სერვისი"             else [hit["oqmi"], hit["moxaz"]]
        res = directory.resolve(names)
        for pos, (name, u) in enumerate(res.surveyors):
            db.add(ItemPerson(item_id=iid, kind=D.P_SURVEYOR, user_id=u.id if u else None,
                              name=None if u else name, position=pos, is_primary=pos == 0,
                              source="regfieldall"))
        if res.surveyors:
            item = db.get(QcItem, iid)
            item.protocol_creator_name = item.protocol_creator_name or res.surveyors[0][0]
            n += 1
    db.flush()
    return n


def import_control(db, raw_dir: Path = RAW_DIR / "control"):
    manifest = json.loads((raw_dir / "manifest.json").read_text(encoding="utf-8"))
    ppl = People(db)
    aliases = load_planner_aliases(db, ppl)

    # ── 1. წაკითხვა ──
    groups = defaultdict(list)
    per_tab = {}
    for t in manifest["tabs"]:
        data = json.loads((raw_dir / t["file"]).read_text(encoding="utf-8"))
        header = data["header"]
        kept = 0
        for i, row in enumerate(data["rows"], start=2):
            rec = row_to_rec(header, row)
            proto, cad = rec.get("proto", "").strip(), rec.get("cad", "").strip()
            if not proto and not cad:
                continue
            rec["_tab"], rec["_row"] = t["name"], i
            groups[(proto, cad)].append(rec)
            kept += 1
        per_tab[t["name"]] = {"rows": len(data["rows"]), "records": kept}

    # ── 2. დუბლიკატები: ერთი ჩანაწერი; სპორადულში თუა — სპორადული რჩება ──
    chosen, duplicates, sporadic_won = [], [], 0
    for key, recs in groups.items():
        spor = [r for r in recs if r["_tab"] == SPORADIC_TAB]
        pool = sorted(spor or recs, key=_score, reverse=True)
        keep = pool[0]
        chosen.append(keep)
        if len(recs) > 1:
            if spor and len(spor) < len(recs):
                sporadic_won += 1
            duplicates.append({"protocol": key[0], "cadcode": key[1],
                               "kept": f"{keep['_tab']}:{keep['_row']}",
                               "dropped": [f"{r['_tab']}:{r['_row']}" for r in recs if r is not keep]})

    # ── 3. ჩაწერა ──
    db.execute(delete(QcItemCode).where(QcItemCode.item_id.in_(
        select(QcItem.id).where(QcItem.is_legacy.is_(True)))))
    db.execute(delete(QcEvent).where(QcEvent.item_id.in_(
        select(QcItem.id).where(QcItem.is_legacy.is_(True)))))
    for model in (ItemPerson, QcReview):
        db.execute(delete(model).where(model.item_id.in_(
            select(QcItem.id).where(QcItem.is_legacy.is_(True)))))
    db.execute(delete(QcItem).where(QcItem.is_legacy.is_(True)))
    db.flush()

    states, verdicts, unmatched = Counter(), Counter(), Counter()
    linked, marks, bad_protocols = Counter(), Counter(), []
    for r in chosen:
        sporadic = r["_tab"] == SPORADIC_TAB
        op_name = clean_name(r.get("qc")) or (None if sporadic else r["_tab"])
        operator = ppl.ensure(op_name, D.QC_OPERATOR, source="control_sheet") if op_name else None
        mgr_name = clean_name(r.get("manager"))
        manager = ppl.get(mgr_name) if mgr_name else None
        if mgr_name and manager is None:
            manager = ppl.ensure(mgr_name, D.FIELD_MANAGER, source="control_sheet", active=False)

        # ამზომველები F და G სვეტებიდან — ყველა სახელი (ზოგჯერ რამდენიმე უჯრაშია)
        names, seen = [], set()
        for nm in split_names(r.get("p1")) + split_names(r.get("p2")):
            k = name_key(nm)
            if k and k not in seen:
                seen.add(k)
                names.append(nm)
        people = []
        for pos, nm in enumerate(names):
            u = ppl.get(nm)
            linked["linked" if u else "unlinked"] += 1
            if u is None:
                unmatched[nm] += 1
            people.append(ItemPerson(kind=D.P_SURVEYOR, user_id=u.id if u else None,
                                     name=None if u else nm, position=pos, is_primary=pos == 0,
                                     source="legacy"))
        if manager is not None:
            people.append(ItemPerson(kind=D.P_MANAGER, user_id=manager.id, position=0,
                                     is_primary=True, source="legacy"))

        state, reason, verdict = derive_state(r)
        states[state] += 1
        verdicts[verdict or "—"] += 1
        added = parse_dt(r.get("added"))
        ldate = parse_dt(r.get("ldate"))
        stime = parse_dt(r.get("stime"))
        resp = D.RESPONSE_FROM_SHEET.get(r.get("jstatus", "").strip())
        fixed = parse_bool(r.get("fixed"))
        cad = r.get("cad", "").strip()
        proto = r.get("proto", "").strip()
        if not re.fullmatch(r"\d{4,9}", proto):
            bad_protocols.append({"tab": r["_tab"], "row": r["_row"], "cadcode": cad})
        answered = qc_answered_last(r)
        returned = (r.get("returned") or "").strip() or None
        extra = dict(r["extra"])
        for k in ("type", "anot", "isdefect", "psync", "rsync", "qtime", "stime", "ldate"):
            if r.get(k, "").strip():
                extra[k] = r[k]

        item = QcItem(
            protocol_number=proto, cadcode=cad,
            kind=D.SPORADIC if sporadic else D.SYSTEMATIC,
            block_key=block_key(cad), zone=zone_of(cad), sector=sector_of(cad),
            drawing_date=parse_dt(r.get("drawing")),
            operator_id=operator.id if operator else None,
            assigned_at=parse_dt(r.get("drawing")) or added,
            field_manager_id=manager.id if manager else None,
            verdict=verdict, original_verdict=verdict, qc_comment=r.get("comment") or None,
            # D — სტატუსის მინიჭების დრო; ძველში გაგზავნამდე იწერებოდა, ამიტომ = პირველი გაგზავნა
            checked_at=added if verdict else None, first_sent_at=added if verdict else None,
            sent_at=added if verdict else None,
            state=state, close_reason=reason,
            closed_at=(max(x for x in (added, ldate, stime) if x) if state == D.CLOSED and
                       any((added, ldate, stime)) else None),
            field_response=resp, field_comment=r.get("kcomment") or None,
            field_responded_at=ldate or (stime if resp else None),
            corrected_confirmed=fixed and verdict in (D.NEGATIVE, D.RECOMMENDATION),
            returned_comment=returned,
            return_count=1 if returned else 0,
            false_fix_count=1 if returned and resp == D.R_CORRECTED else 0,
            dispute_count=1 if resp == D.R_DISAGREE else 0,
            qcm_explanation=returned if resp == D.R_DISAGREE else None,
            qcm_decision=D.QCM_STAND if resp == D.R_DISAGREE and answered else None,
            # ძველი დაბრუნებები მენეჯერთან იყო — პასუხობს მენეჯერი (პასუხი №15)
            manager_required=state == D.RETURNED_TO_MANAGER,
            is_legacy=True, legacy_source=r["_tab"], legacy_row=r["_row"],
            legacy_extra=extra or None,
        )
        item.people = people
        db.add(item)
        db.flush()
        for code in parse_codes(r.get("codes")):
            db.add(QcItemCode(item_id=item.id, code=code, source="legacy"))
        mark = D.LEGACY_TYPE_TO_MARK.get((r.get("type") or "").strip().lower())
        if mark:                                   # V სვეტი → QC მენეჯერის შეფასება (შეზღუდული)
            marks[mark] += 1
            db.add(QcReview(item_id=item.id, marks=[mark], comment=(r.get("anot") or "").strip() or None,
                            source="legacy", at=added or dt.datetime.now()))
        snap = {k: r.get(k) for k in ("added", "status", "jstatus", "ldate", "fixed", "returned",
                                     "qtime", "stime", "manager", "type") if r.get(k)}
        db.add(QcEvent(item_id=item.id, at=added or item.assigned_at or dt.datetime.now(),
                       action="legacy_import", to_state=state, approx=True,
                       comment=f"ძველი ცხრილი — {r['_tab']}, მწკრივი {r['_row']}",
                       data={"snapshot": snap}))

    db.flush()
    filled = fill_people_from_source(db)

    report = {
        "fetched_at": manifest.get("fetched_at"),
        "surveyors_from_regfieldall": filled,
        "per_tab": per_tab,
        "source_rows": sum(v["rows"] for v in per_tab.values()),
        "records": sum(len(v) for v in groups.values()),
        "unique_items": len(chosen),
        "duplicates": len(duplicates),
        "duplicates_sporadic_kept": sporadic_won,
        "duplicate_samples": duplicates[:50],
        "legacy_review_marks": dict(marks),
        "bad_protocol_numbers": bad_protocols,
        "states": dict(states), "verdicts": dict(verdicts),
        "planner_links": dict(linked),
        "unmatched_planners": unmatched.most_common(60),
        "unmatched_planner_names": len(unmatched),
        "aliases_loaded": aliases,
    }
    db.add(ImportRun(kind="legacy_control", source={"fetched_at": manifest.get("fetched_at"),
                                                   "tabs": [t["name"] for t in manifest["tabs"]]},
                     report=report, item_count=len(chosen)))
    db.flush()
    return report
