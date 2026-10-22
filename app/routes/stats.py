# -*- coding: utf-8 -*-
"""სტატისტიკა — როლის მიხედვით განსხვავებული განყოფილებები, ერთიანი გამოთვლით."""
import datetime as dt
import io

from fastapi import APIRouter, Depends, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from .. import domain as D
from .. import scope
from ..db import get_db
from ..models import DefectCode, Mission
from ..security import current
from ..services import stats as S
from .common import render
from .items import zone_names

router = APIRouter()

SECTIONS = {
    D.COORDINATOR: ["kpi", "trend", "operators", "managers", "surveyors", "codes", "zones"],
    # სამსახურის უფროსი — ხარვეზების სტატისტიკა, „კონტროლის სტატისტიკის" (ოპერატორების) გარეშე
    D.HEAD_OF_SERVICE: ["kpi", "trend", "managers", "surveyors", "codes", "zones"],
    D.ADMIN: ["kpi", "trend", "operators", "managers", "surveyors", "codes", "zones", "restricted"],
    D.QC_MANAGER: ["kpi", "trend", "operators", "managers", "codes", "zones", "restricted"],
    D.QC_OPERATOR: ["kpi", "trend", "operators", "codes", "zones"],
    D.FIELD_MANAGER: ["kpi", "trend", "surveyors", "codes", "zones", "restricted"],
    D.SURVEYOR: ["kpi", "trend", "codes", "zones"],
}


def compute(db, ctx, f):
    sec = SECTIONS.get(ctx.role, ["kpi"])
    out = {"role": ctx.role, "sections": sec}
    if "kpi" in sec:
        out["kpi"] = S.kpi(db, ctx, f)
    if "trend" in sec:
        out["trend"] = S.trend(db, ctx, f)
    if "operators" in sec:
        out["operators"] = S.by_operator(db, ctx, f)
    if "managers" in sec:
        out["managers"] = S.by_manager(db, ctx, f)
    if "surveyors" in sec:
        out["surveyors"] = S.by_surveyor(db, ctx, f)
    if "codes" in sec:
        out["codes"] = S.codes(db, ctx, f)
    if "zones" in sec:
        out["zones"] = S.zones(db, ctx, f, zone_names())
    if "restricted" in sec:
        out["restricted"] = S.restricted(db, ctx, f)
    return out


@router.get("/stats")
def stats_page(request: Request, ctx=Depends(current), db=Depends(get_db)):
    opts = S.filter_options(db, ctx)
    missions = db.scalars(select(Mission).order_by(Mission.n)).all()
    codes = db.scalars(select(DefectCode).order_by(DefectCode.sort)).all()
    groups = sorted({c.group for c in codes if c.severity != "pending"})
    surveyors = []
    if ctx.role == D.FIELD_MANAGER:
        surveyors = [(u.id, u.full_name) for u in S.team_surveyors(db, ctx.user.id)]
    return render(request, ctx, "stats/index.html", active="stats", opts=opts, missions=missions,
                  codes=codes, groups=groups, surveyors=surveyors,
                  restricted=scope.can_see_restricted_stats(ctx),
                  false_fix=scope.can_see_false_fix(ctx), today=dt.date.today())


@router.get("/stats/tabler-preview")
def stats_tabler_preview(request: Request, ctx=Depends(current)):
    """საცდელი დიზაინი (Tabler). ძველ /stats გვერდს არ ცვლის; იგივე /api/stats მონაცემი."""
    return render(request, ctx, "stats/tabler_preview.html")


@router.get("/api/stats")
def stats_api(request: Request, ctx=Depends(current), db=Depends(get_db)):
    f = S.parse_filters(request.query_params)
    return compute(db, ctx, f)


@router.get("/stats/export.xlsx")
def stats_export(request: Request, ctx=Depends(current), db=Depends(get_db)):
    f = S.parse_filters(request.query_params)
    data = compute(db, ctx, f)
    wb = Workbook()
    ws = wb.active
    ws.title = "ჯამური"
    k = data.get("kpi", {})
    labels = [("checked", "შემოწმებული"), ("positive", "დადებითი (+დაზუსტება)"),
              ("negative", "უარყოფითი"), ("recommendation", "რეკომენდაცია"),
              ("defect_rate", "ხარვეზი %"), ("critical_rate", "კრიტიკული %"),
              ("corrected", "გასწორდა"), ("corrected_rate", "გასწორდა %"),
              ("in_field", "ველზე ღია"), ("awaiting_qc", "QC-ის გადასამოწმებელი"),
              ("pending", "გასავლელი"), ("response_median", "პასუხის მედიანა (დღე)")]
    ws.append(["მაჩვენებელი", "მნიშვნელობა"])
    for key, lbl in labels:
        ws.append([lbl, k.get(key)])
    tables = {"operators": ("ოპერატორები", "name"), "managers": ("მენეჯერები", "name"),
              "surveyors": ("ამზომველები", "name"), "zones": ("ზონები", "name")}
    cols = ["checked", "positive", "negative", "recommendation", "defect_rate", "critical_rate",
            "corrected", "corrected_rate", "in_field", "awaiting_qc", "disputed"]
    heads = ["შემოწმებული", "დადებითი", "უარყოფითი", "რეკომენდაცია", "ხარვეზი %", "კრიტიკული %",
             "გასწორდა", "გასწორდა %", "ველზე ღია", "გადასამოწმებელი", "დავა"]
    for key, (title, name) in tables.items():
        if key not in data:
            continue
        w = wb.create_sheet(title)
        w.append(["სახელი"] + heads)
        for r in data[key]:
            w.append([r.get(name)] + [r.get(c) for c in cols])
    if "codes" in data:
        w = wb.create_sheet("კოდები")
        w.append(["კოდი", "დასახელება", "ჯგუფი", "სიმძიმე", "უარყოფითში", "რეკომენდაციაში"])
        for c in data["codes"]:
            w.append([c["code"], c["name"], c["group"], c["severity"], c["negative"], c["recommendation"]])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        sh.column_dimensions["A"].width = 32
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition":
                                      f'attachment; filename="statistika_{dt.date.today()}.xlsx"'})
