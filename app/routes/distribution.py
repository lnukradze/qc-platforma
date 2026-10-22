# -*- coding: utf-8 -*-
"""
განაწილება — უშვებს QC მენეჯერი ან ადმინისტრატორი.

  /distribution            სისტემური: პერიოდი, პროცენტები (ზონა/სექტორი/მენეჯერი), ოპერატორები
  /distribution/sporadic   სპორადული: IT-ის ყოველდღიური სიის ატვირთვა → ფილტრი → განაწილება
  /distribution/unsorted   „დაუჯგუფებელი": წასული ამზომველების ოქმები — მენეჯერის ხელით მინიჭება
"""
import datetime as dt
import io
import json
import tempfile
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from .. import domain as D
from ..config import APP_DIR, DATA_DIR
from ..db import get_db
from ..models import (DistributionBacklog, DistributionPreset, DistributionRun, ManagerOverride,
                      Setting, SporadicImport, User, UserRole)
from ..security import audit, require, same_origin
from ..services import sporadic
from ..services.distribution import service, source
from ..services.people import clean_name, name_key
from .common import redirect, render
from .items import zone_names

router = APIRouter()
ROLES = (D.QC_MANAGER, D.ADMIN)
OPEN_STATUSES = ["შედგენილია -საველე განცხადება მიღებულია", "აღწერის პროცესი-საველე განცხადება",
                 "უარი საველე წარმოებაზე"]
MAX_SPORADIC_MB = 40
_SECTORS = []


def sector_options():
    """[(„71.56", „71.56 — ბაზალეთი (დუშეთი)")] — პროცენტის წესისთვის."""
    if not _SECTORS:
        d = json.loads((APP_DIR / "static" / "geo" / "names.json").read_text(encoding="utf-8"))
        zones = d.get("zones", {})
        for z, secs in sorted(d.get("sectors", {}).items()):
            for s, nm in sorted(secs.items()):
                if nm:
                    _SECTORS.append((f"{z}.{s}", f"{z}.{s} — {nm} ({zones.get(z, z)})"))
    return _SECTORS


def _operators(db):
    return db.scalars(select(User).join(UserRole).where(UserRole.role == D.QC_OPERATOR,
                                                       User.active.is_(True))
                      .order_by(User.full_name)).all()


def _managers(db):
    return db.execute(select(User.id, User.full_name).join(UserRole)
                      .where(UserRole.role == D.FIELD_MANAGER, User.active.is_(True))
                      .order_by(User.full_name)).all()


def _ops_from_form(form):
    ops = {}
    for k in form.keys():
        if k.startswith("op_on_"):
            oid = k[6:]
            raw = str(form.get(f"op_q_{oid}") or "").strip()
            ops[oid] = int(raw) if raw.isdigit() else 0
    return ops


def _params_from_form(form) -> dict:
    rules = []
    for t, k, p in zip(form.getlist("rule_type"), form.getlist("rule_key"), form.getlist("rule_pct")):
        k = str(k or "").split(" ")[0].strip()
        if t and k and str(p).strip():
            rules.append({"type": t, "key": k, "pct": float(p)})
    return {
        "date_from": form.get("date_from"), "date_to": form.get("date_to") or form.get("date_from"),
        "operators": {o: q for o, q in _ops_from_form(form).items() if q > 0},
        "default_pct": float(form.get("default_pct") or 70),
        "rules": rules,
        "exclude_zones": form.getlist("exclude_zones"),
        "exclude_statuses": form.getlist("exclude_statuses"),
        "seed": int(form.get("seed")) if (form.get("seed") or "").isdigit() else None,
    }


def _subnav():
    return [("/distribution", "სისტემური", "system"), ("/distribution/sporadic", "სპორადული", "sporadic"),
            ("/distribution/unsorted", "დაუჯგუფებელი", "unsorted")]


@router.get("/distribution")
def page(request: Request, preset: int | None = None, ctx=Depends(require(*ROLES)),
         db=Depends(get_db)):
    d = service.defaults(db)
    src = source.source_info()
    today = dt.date.today()
    max_date = dt.date.fromisoformat(src["max_date"]) if src.get("ok") and src.get("max_date") else today
    form = {"date_from": (max_date - dt.timedelta(days=6)).isoformat(),
            "date_to": max_date.isoformat(), "default_pct": d.get("default_pct", 70),
            "rules": [], "exclude_statuses": d.get("exclude_statuses", []),
            "per_operator": d.get("per_operator", 450), "operators": {}, "exclude_zones": []}
    presets = db.scalars(select(DistributionPreset).where(DistributionPreset.active.is_(True))
                         .order_by(DistributionPreset.id.desc())).all()
    if preset:
        pr = db.get(DistributionPreset, preset)
        if pr:
            form.update({k: v for k, v in pr.params.items() if v is not None})
    runs = db.scalars(select(DistributionRun).options(selectinload(DistributionRun.created_by))
                      .order_by(DistributionRun.id.desc()).limit(12)).all()
    backlog = db.query(DistributionBacklog).filter(DistributionBacklog.status == "open").count()
    return render(request, ctx, "distribution/index.html", active="distribution", sub="system",
                  subnav=_subnav(), form=form, operators=_operators(db), managers=_managers(db),
                  zones=sorted(zone_names().items()), sectors=sector_options(), source=src,
                  presets=presets, runs=runs, open_statuses=OPEN_STATUSES, backlog=backlog)


@router.post("/distribution/preview")
async def preview(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    form = await request.form()
    params = _params_from_form(form)
    try:
        pv = service.prepare(db, params)
        err = None
    except (ValueError, FileNotFoundError) as e:
        pv, err = None, str(e)
    return render(request, ctx, "distribution/_preview.html", pv=pv, err=err, zone_names=zone_names())


@router.post("/distribution/confirm")
async def confirm(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    form = await request.form()
    try:
        run = service.confirm(db, form.get("token", ""), ctx.user, ctx.role)
    except ValueError as e:
        return redirect("/distribution", err=str(e))
    audit(db, ctx.user, "distribution_confirmed", "distribution_run", run.id,
          {"items": run.item_count}, request)
    db.commit()
    carried = (run.summary or {}).get("carried_over") or 0
    msg = f"განაწილდა {run.item_count} ოქმი — ოპერატორებს უკვე უჩანთ „გასავლელში“."
    if carried:
        msg += f" {carried} ოქმი ვერ ჩაეტია და შემდეგ განაწილებაზე პირველი შევა."
    return redirect(f"/distribution/runs/{run.id}", ok=msg)


@router.get("/distribution/runs/{run_id}")
def run_detail(run_id: int, request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    run = db.get(DistributionRun, run_id)
    if run is None:
        raise HTTPException(404, "განაწილება ვერ მოიძებნა.")
    return render(request, ctx, "distribution/run.html", active="distribution", sub="system",
                  subnav=_subnav(), run=run)


@router.post("/distribution/presets")
async def save_preset(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    form = await request.form()
    params = _params_from_form(form)
    name = (form.get("preset_name") or "").strip() or f"წესი {dt.date.today():%d.%m}"
    vf = form.get("valid_from") or None
    vt = form.get("valid_to") or None
    p = DistributionPreset(name=name, params=params, created_by_id=ctx.user.id,
                           valid_from=dt.date.fromisoformat(vf) if vf else None,
                           valid_to=dt.date.fromisoformat(vt) if vt else None)
    db.add(p)
    db.commit()
    if request.headers.get("hx-request"):
        from fastapi.responses import HTMLResponse
        from markupsafe import escape
        return HTMLResponse(f'<div class="notice ok">✓ წესი „{escape(name)}“ შენახულია — '
                            f'<a href="/distribution?preset={p.id}">ჩატვირთვა</a></div>')
    return redirect(f"/distribution?preset={p.id}", ok=f"წესი „{name}“ შენახულია.")


@router.get("/distribution/export/{token}")
def export(token: str, ctx=Depends(require(*ROLES))):
    pv = service.get_preview(token) or sporadic.get_preview(token)
    if pv is None:
        raise HTTPException(404, "გადახედვის ვადა ამოიწურა.")
    wb = Workbook()
    ws = wb.active
    ws.title = "განაწილება"
    cols = ["ნახაზის შეტანა / დათვალიერება", "საკადასტრო", "ბლოკი", "ოქმის N", "ოქმის შემდგენი",
            "ამზომველები", "მენეჯერი", "ზონა", "პროცენტი", "წინა განაწილებიდან", "კონტროლის სპეციალისტი"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    for r in pv["rows"]:
        ws.append([r.get("disp_dt") or r.get("survey_day"), r["cadcode"], r.get("block", ""),
                   r["protocol"], r.get("creator"), ", ".join(n for n, _ in r.get("surveyors") or []),
                   r.get("manager"), r.get("zone"), r.get("pct", ""),
                   "დიახ" if r.get("backlog") else "", r.get("operator")])
    for i, w in enumerate([22, 24, 20, 12, 26, 34, 26, 8, 10, 12, 26], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"ganawileba_{dt.datetime.now():%Y-%m-%d_%H-%M}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fn}"'})


# ── სპორადული ─────────────────────────────────────────────────────────
@router.get("/distribution/sporadic")
def sporadic_page(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    s = db.get(Setting, "sporadic_defaults")
    defaults = dict(s.value) if s else {}
    imports = db.scalars(select(SporadicImport).options(selectinload(SporadicImport.created_by))
                         .order_by(SporadicImport.id.desc()).limit(10)).all()
    return render(request, ctx, "distribution/sporadic.html", active="distribution", sub="sporadic",
                  subnav=_subnav(), defaults=defaults, operators=_operators(db), imports=imports,
                  today=dt.date.today())


@router.post("/distribution/sporadic/preview")
async def sporadic_preview(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    form = await request.form()
    up = form.get("file")
    if up is None or not getattr(up, "filename", ""):
        return render(request, ctx, "distribution/_sporadic_preview.html", pv=None,
                      err="აირჩიეთ ფაილი (xlsx).")
    content = await up.read()
    if len(content) > MAX_SPORADIC_MB * 1024 * 1024:
        return render(request, ctx, "distribution/_sporadic_preview.html", pv=None,
                      err=f"ფაილი {MAX_SPORADIC_MB} MB-ზე დიდია.")
    # 🔐 ფაილში პირადი მონაცემია — დროებით ვინახავთ, ვკითხულობთ საჭირო სვეტებს და მაშინვე ვშლით
    tmp_dir = DATA_DIR / "tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=tmp_dir)
    try:
        with open(fd, "wb") as f:
            f.write(content)
        rows = sporadic.read_file(tmp)
    except (ValueError, OSError, KeyError) as e:
        return render(request, ctx, "distribution/_sporadic_preview.html", pv=None,
                      err=f"ფაილი ვერ წავიკითხე: {e}")
    finally:
        Path(tmp).unlink(missing_ok=True)
    params = {
        "survey_day_mode": form.get("survey_day_mode") or "upto_yesterday",
        "lookback_days": int(form.get("lookback_days") or 7),
        "call_status_completed_only": form.get("call_status_completed_only") == "on",
        "exclude_open_status": form.get("exclude_open_status") == "on",
        "operators": _ops_from_form(form),
        "seed": int(form.get("seed")) if (form.get("seed") or "").isdigit() else None,
        "filename": up.filename,
    }
    today = form.get("today")
    try:
        pv = sporadic.prepare(db, rows, params,
                              today=dt.date.fromisoformat(today) if today else None)
        err = None
    except (ValueError, FileNotFoundError) as e:
        pv, err = None, str(e)
    return render(request, ctx, "distribution/_sporadic_preview.html", pv=pv, err=err)


@router.post("/distribution/sporadic/confirm")
async def sporadic_confirm(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    form = await request.form()
    pv = sporadic.get_preview(form.get("token", ""))
    try:
        imp = sporadic.confirm(db, form.get("token", ""), ctx.user, ctx.role,
                               filename=(pv or {}).get("params", {}).get("filename"))
    except ValueError as e:
        return redirect("/distribution/sporadic", err=str(e))
    audit(db, ctx.user, "sporadic_distributed", "sporadic_import", imp.id,
          {"items": imp.item_count}, request)
    db.commit()
    return redirect("/distribution/sporadic", ok=f"სპორადული: განაწილდა {imp.item_count} ოქმი.")


# ── დაუჯგუფებელი (მენეჯერის ხელით მინიჭება) ─────────────────────────────
@router.get("/distribution/unsorted")
def unsorted(request: Request, date_from: str = "", date_to: str = "",
             ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    src = source.source_info()
    today = dt.date.today()
    max_date = dt.date.fromisoformat(src["max_date"]) if src.get("ok") and src.get("max_date") else today
    date_from = date_from or (max_date - dt.timedelta(days=30)).isoformat()
    date_to = date_to or max_date.isoformat()
    people, total, err = [], 0, None
    try:
        people, total = service.unsorted_report(db, date_from, date_to,
                                                service.defaults(db).get("exclude_statuses"))
    except FileNotFoundError as e:
        err = str(e)
    overrides = db.scalars(select(ManagerOverride).options(selectinload(ManagerOverride.manager),
                                                           selectinload(ManagerOverride.set_by))
                           .order_by(ManagerOverride.set_at.desc())).all()
    return render(request, ctx, "distribution/unsorted.html", active="distribution", sub="unsorted",
                  subnav=_subnav(), people=people, total=total, err=err, date_from=date_from,
                  date_to=date_to, managers=_managers(db), overrides=overrides)


@router.post("/distribution/unsorted/assign")
async def unsorted_assign(request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    form = await request.form()
    name = clean_name(form.get("person_name"))
    mid = str(form.get("manager_id") or "")
    if not name or not mid.isdigit():
        return redirect("/distribution/unsorted", err="აირჩიეთ მენეჯერი.")
    k = name_key(name)
    o = db.scalar(select(ManagerOverride).where(ManagerOverride.person_key == k))
    if o is None:
        o = ManagerOverride(person_key=k, person_name=name, manager_id=int(mid))
        db.add(o)
    o.manager_id, o.set_by_id, o.set_at = int(mid), ctx.user.id, dt.datetime.now()
    o.note = (form.get("note") or "").strip() or None
    audit(db, ctx.user, "manager_override", "person", None, {"person": name, "manager": int(mid)}, request)
    db.commit()
    back = form.get("back") or "/distribution/unsorted"
    return redirect(back, ok=f"„{name}“ — მენეჯერი მიეწერა. მისი ოქმები შემდეგ განაწილებაში ჩვეულებრივ შევა.")


@router.post("/distribution/overrides/{oid}/delete")
async def override_delete(oid: int, request: Request, ctx=Depends(require(*ROLES)), db=Depends(get_db)):
    same_origin(request)
    o = db.get(ManagerOverride, oid)
    if o is not None:
        audit(db, ctx.user, "manager_override_removed", "person", None, {"person": o.person_name}, request)
        db.delete(o)
        db.commit()
    return redirect("/distribution/unsorted", ok="მინიჭება წაიშალა.")
